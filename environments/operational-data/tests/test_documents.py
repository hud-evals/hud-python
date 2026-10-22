"""Deliverables render as text the judge can read, and unreadable ones say why."""

from pathlib import Path

import pytest
from docx import Document
from openpyxl import Workbook
from pptx import Presentation
from pptx.util import Inches
from pypdf import PdfWriter

import agent_workspace
from documents import UnreadableDocument, document_text


def test_a_workbook_shows_every_sheet_and_unevaluated_formulas(tmp_path: Path):
    path = tmp_path / "run.xlsx"
    workbook = Workbook()
    run = workbook.active
    assert run is not None
    run.title = "Run"
    run.append(["Bill", "Pay"])
    run.append(["ML-4471", 44900])
    run.append([])
    run.append(["Total", "=SUM(B2:B2)"])
    workbook.create_sheet("Notes").append(["Short-pay per CM-0219"])
    workbook.save(path)

    assert document_text(path) == (
        '## Sheet: Run\n["Bill", "Pay"]\n["ML-4471", 44900]\n'
        '["Total", "[uncached formula: =SUM(B2:B2)]"]\n'
        '## Sheet: Notes\n["Short-pay per CM-0219"]'
    )


def test_a_word_document_shows_paragraphs_and_tables(tmp_path: Path):
    path = tmp_path / "memo.docx"
    document = Document()
    document.add_paragraph("Payment run for April 1 to 10.")
    table = document.add_table(rows=1, cols=2)
    table.rows[0].cells[0].text = "ML-4471"
    table.rows[0].cells[1].text = "44,900.00"
    document.save(str(path))

    assert document_text(path) == "Payment run for April 1 to 10.\n## Table\nML-4471\t44,900.00"


def test_a_presentation_shows_each_slide(tmp_path: Path):
    path = tmp_path / "deck.pptx"
    deck = Presentation()
    slide = deck.slides.add_slide(deck.slide_layouts[6])
    slide.shapes.add_textbox(Inches(1), Inches(1), Inches(4), Inches(1)).text_frame.text = "Run"
    deck.save(str(path))

    assert document_text(path) == "## Slide 1\nRun"


def test_a_pdf_shows_each_page(tmp_path: Path):
    path = tmp_path / "run.pdf"
    writer = PdfWriter()
    writer.add_blank_page(width=612, height=792)
    with path.open("wb") as out:
        writer.write(out)

    assert document_text(path) == "## Page 1\n"


def test_text_files_are_read_as_is(tmp_path: Path):
    path = tmp_path / "run.csv"
    path.write_text("bill,pay\nML-4471,44900\n")
    assert document_text(path) == "bill,pay\nML-4471,44900\n"


@pytest.mark.parametrize(
    ("name", "content", "reason"),
    [
        ("run.xlsx", b"not a workbook", "could not be read as .xlsx"),
        ("run.bin", b"\xff\xfe\x00\x81", "not a text file"),
    ],
)
def test_unreadable_files_say_why(tmp_path: Path, name: str, content: bytes, reason: str):
    path = tmp_path / name
    path.write_bytes(content)
    with pytest.raises(UnreadableDocument, match=reason):
        document_text(path)


def test_deliverables_report_their_status(tmp_path: Path):
    outputs = tmp_path / "outputs"
    outputs.mkdir()
    (outputs / "notes.md").write_text("x" * 40_000)
    (outputs / "broken.docx").write_bytes(b"not a document")
    secret = tmp_path / "secret.txt"
    secret.write_text("HUD_API_KEY=...")
    (outputs / "run.md").symlink_to(secret)
    (outputs / "linked").symlink_to(tmp_path)
    saved, broken, missing, link, through_link = (
        agent_workspace.read_deliverable(tmp_path, target)
        for target in agent_workspace.deliverable_targets(
            tmp_path, ["notes.md", "broken.docx", "run.xlsx", "run.md", "linked/secret.txt"]
        )
    )

    assert (saved.path, saved.status) == ("outputs/notes.md", "saved")
    assert saved.text == "x" * 40_000
    assert (broken.status, missing.status) == ("unreadable", "missing")
    assert broken.text.startswith("The file could not be read as .docx")
    assert (link.status, through_link.status) == ("unreadable", "unreadable")
    assert "HUD_API_KEY" not in link.text + through_link.text
