"""The sample exports load, and the facts the example tasks hinge on are reachable through tools."""

from datetime import date, datetime
from zoneinfo import ZoneInfo

import pytest
from exports import SAMPLES, load
from openpyxl import load_workbook

from data_sources.base import DataSource, TaskContext
from data_sources.sources import SOURCES, Calendar, Jira, Slack

PACIFIC = ZoneInfo("America/Los_Angeles")


@pytest.mark.parametrize("source", list(SOURCES.values()), ids=list(SOURCES))
def test_every_sample_export_loads(source: type[DataSource], context: TaskContext):
    loaded = load(source, context)
    names = [item.name for item in loaded.tools()]
    assert loaded.describe()
    assert names
    assert all(name.startswith(f"{source.tool_prefix()}_") for name in names)


def test_board_review_has_exactly_one_common_hour(context: TaskContext):
    calendar = load(Calendar, context)
    slots = calendar.find_free_time(
        ["jordan", "dana", "sam", "elena"], date(2026, 4, 6), date(2026, 4, 10), 60
    )
    assert [(slot.start, slot.end) for slot in slots] == [
        (datetime(2026, 4, 9, 14, tzinfo=PACIFIC), datetime(2026, 4, 9, 15, tzinfo=PACIFIC))
    ]


def test_march_close_status_comes_from_the_close_sprint(context: TaskContext):
    issues = load(Jira, context).search_issues('sprint = "March 2026 Close"').items
    assert {issue.key: issue.status for issue in issues} == {
        "CLOSE-11": "To Do",
        "CLOSE-12": "To Do",
        "CLOSE-13": "In Progress",
        "CLOSE-14": "To Do",
        "CLOSE-15": "In Progress",
        "CLOSE-16": "Done",
        "CLOSE-17": "To Do",
        "CLOSE-18": "To Do",
    }
    flash = load(Calendar, context).list_events(
        start=date(2026, 4, 1), end=date(2026, 4, 1), query="Flash"
    )
    assert {event.start for event in flash.items} == {datetime(2026, 4, 1, 15, tzinfo=PACIFIC)}


def test_meridian_decision_is_in_slack_and_jira(context: TaskContext):
    found = load(Slack, context).search_messages("short-pay").items
    assert "ap-payables" in {message.channel for message in found}
    ticket = load(Jira, context).get_issue("CLOSE-13")
    assert ticket.due == date(2026, 4, 4)
    assert any("CM-0219" in comment.body for comment in ticket.comments)


def test_payment_run_turns_on_the_aging_slack_and_jira(context: TaskContext):
    sheet = load_workbook(SAMPLES / "attachments" / "AP aging 2026-03-31.xlsx").active
    assert sheet is not None
    header = next(row for row in sheet.iter_rows(values_only=True) if row[0] == "Vendor")
    bills = [
        dict(zip(header, row, strict=True))
        for row in sheet.iter_rows(values_only=True)
        if isinstance(row[1], str) and "-" in row[1]
    ]
    due = {
        bill["Bill No."]: bill["Open Balance"]
        for bill in bills
        if bill["Due Date"] <= datetime(2026, 4, 10)
    }
    assert due == {
        "PT-0471": 2_300,
        "ML-4471": 48_200,
        "KP-2026-02": 18_000,
        "HL-55102": 1_980,
        "PT-0488": 3_150,
        "NF-2291": 27_640,
    }
    paid = load(Slack, context).search_messages("check 1047").items
    assert any("PT-0471" in message.text for message in paid)
    ticket = load(Jira, context).get_issue("CLOSE-13")
    assert any("$44,900" in comment.body for comment in ticket.comments)
    assert sum(due.values()) - 2_300 - 48_200 + 44_900 == 95_670


def test_summit_reship_is_tracked_in_ops_88(context: TaskContext):
    ticket = load(Jira, context).get_issue("OPS-88")
    assert ticket.status == "In Progress"
    assert any("Harborline" in comment.body for comment in ticket.comments)
    found = load(Slack, context).search_messages("INV-2026-0134").items
    assert "sales-pipeline" in {message.channel for message in found}
