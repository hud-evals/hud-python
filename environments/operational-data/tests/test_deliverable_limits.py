"""Privileged deliverable reads stay within regular files and bounded document work."""

from __future__ import annotations

import os
from io import BytesIO
from pathlib import Path
from xml.etree import ElementTree
from zipfile import ZIP_DEFLATED, ZipFile

import pytest
from openpyxl import Workbook

import agent_workspace
import documents


@pytest.fixture
def deliverable_root(tmp_path: Path) -> Path:
    root = tmp_path / "workspace"
    (root / "outputs").mkdir(parents=True)
    return root


def _workbook(path: Path) -> None:
    workbook = Workbook()
    sheet = workbook.active
    assert sheet is not None
    sheet.append(["Bill", "Pay"])
    sheet.append(["ML-4471", 44900])
    workbook.save(path)
    workbook.close()


def test_regular_text_is_preserved_without_silent_truncation(deliverable_root: Path):
    target = deliverable_root / "outputs" / "notes.md"
    text = "payment run\n" * 4000
    target.write_text(text, encoding="utf-8")

    result = agent_workspace.read_deliverable(deliverable_root, target)

    assert result.path == "outputs/notes.md"
    assert result.status == "saved"
    assert result.text == text


def test_text_at_the_utf8_byte_limit_is_preserved(deliverable_root: Path):
    target = deliverable_root / "outputs" / "notes.md"
    text = "é" * (documents.MAX_TEXT_BYTES // 2)
    target.write_text(text, encoding="utf-8")

    result = agent_workspace.read_deliverable(deliverable_root, target)

    assert result.status == "saved"
    assert result.text == text


def test_regular_workbook_is_read_with_its_original_format(deliverable_root: Path):
    target = deliverable_root / "outputs" / "run.xlsx"
    _workbook(target)

    result = agent_workspace.read_deliverable(deliverable_root, target)

    assert result.status == "saved"
    assert result.text == '## Sheet: Sheet\n["Bill", "Pay"]\n["ML-4471", 44900]'


@pytest.mark.parametrize("name", ["missing.md", "missing/notes.md"])
def test_missing_deliverables_remain_missing(deliverable_root: Path, name: str):
    result = agent_workspace.read_deliverable(deliverable_root, deliverable_root / "outputs" / name)

    assert result.status == "missing"


def test_leaf_symlinks_are_unreadable(deliverable_root: Path):
    secret = deliverable_root.parent / "secret.md"
    secret.write_text("private verifier data", encoding="utf-8")
    target = deliverable_root / "outputs" / "notes.md"
    target.symlink_to(secret)

    result = agent_workspace.read_deliverable(deliverable_root, target)

    assert result.status == "unreadable"
    assert "private verifier data" not in result.text


def test_ancestor_symlinks_are_rejected_even_when_they_stay_inside_outputs(deliverable_root: Path):
    real = deliverable_root / "outputs" / "real"
    real.mkdir()
    (real / "notes.md").write_text("a linked document", encoding="utf-8")
    (deliverable_root / "outputs" / "linked").symlink_to(real, target_is_directory=True)

    result = agent_workspace.read_deliverable(
        deliverable_root, deliverable_root / "outputs" / "linked" / "notes.md"
    )

    assert result.status == "unreadable"


def test_replacing_outputs_with_a_symlink_does_not_move_the_trust_boundary(deliverable_root: Path):
    elsewhere = deliverable_root.parent / "elsewhere"
    elsewhere.mkdir()
    (elsewhere / "notes.md").write_text("outside outputs", encoding="utf-8")
    outputs = deliverable_root / "outputs"
    outputs.rmdir()
    outputs.symlink_to(elsewhere, target_is_directory=True)

    result = agent_workspace.read_deliverable(deliverable_root, outputs / "notes.md")

    assert result.status == "unreadable"
    assert "outside outputs" not in result.text


def test_directories_are_unreadable_deliverables(deliverable_root: Path):
    target = deliverable_root / "outputs" / "notes.md"
    target.mkdir()

    result = agent_workspace.read_deliverable(deliverable_root, target)

    assert result.status == "unreadable"


@pytest.mark.skipif(not hasattr(os, "mkfifo"), reason="requires POSIX named pipes")
def test_named_pipes_are_unreadable_without_waiting_for_a_writer(deliverable_root: Path):
    target = deliverable_root / "outputs" / "notes.md"
    os.mkfifo(target)

    result = agent_workspace.read_deliverable(deliverable_root, target)

    assert result.status == "unreadable"


def test_files_above_the_input_byte_limit_are_unreadable(deliverable_root: Path):
    target = deliverable_root / "outputs" / "notes.bin"
    with target.open("wb") as out:
        out.seek(documents.MAX_FILE_BYTES)
        out.write(b"\0")

    result = agent_workspace.read_deliverable(deliverable_root, target)

    assert result.status == "unreadable"


@pytest.mark.parametrize("character", ["x", "é"])
def test_text_above_the_utf8_byte_limit_is_unreadable(deliverable_root: Path, character: str):
    target = deliverable_root / "outputs" / "notes.md"
    text = character * (documents.MAX_TEXT_BYTES // len(character.encode("utf-8")) + 1)
    target.write_text(text, encoding="utf-8")

    result = agent_workspace.read_deliverable(deliverable_root, target)

    assert result.status == "unreadable"


def test_expanded_ooxml_size_is_checked_before_the_document_parser(
    deliverable_root: Path, monkeypatch: pytest.MonkeyPatch
):
    target = deliverable_root / "outputs" / "run.xlsx"
    _workbook(target)
    with ZipFile(target, "a", compression=ZIP_DEFLATED) as archive:
        with archive.open("padding.bin", "w") as out:
            remaining = documents.MAX_EXPANDED_DOCUMENT_BYTES + 1
            chunk = b"x" * (1 << 20)
            while remaining:
                piece = chunk[:remaining]
                out.write(piece)
                remaining -= len(piece)
    assert target.stat().st_size < documents.MAX_FILE_BYTES

    def parser_must_not_run(*args: object, **kwargs: object):
        pytest.fail("an oversized OOXML archive reached the document parser")

    monkeypatch.setattr(documents, "load_workbook", parser_must_not_run)

    result = agent_workspace.read_deliverable(deliverable_root, target)

    assert result.status == "unreadable"


def test_advertised_huge_workbook_ranges_are_unreadable(deliverable_root: Path):
    target = deliverable_root / "outputs" / "run.xlsx"
    _workbook(target)
    replacement = BytesIO()
    with ZipFile(target) as source, ZipFile(replacement, "w") as output:
        for item in source.infolist():
            data = source.read(item)
            if item.filename == "xl/worksheets/sheet1.xml":
                sheet = ElementTree.fromstring(data)
                dimension = sheet.find(
                    "{http://schemas.openxmlformats.org/spreadsheetml/2006/main}dimension"
                )
                assert dimension is not None
                dimension.set("ref", "A1:XFD1048576")
                data = ElementTree.tostring(sheet)
            output.writestr(item, data)
    target.write_bytes(replacement.getvalue())

    result = agent_workspace.read_deliverable(deliverable_root, target)

    assert result.status == "unreadable"


def test_workbook_cell_budget_accumulates_across_sheets(
    deliverable_root: Path, monkeypatch: pytest.MonkeyPatch
):
    target = deliverable_root / "outputs" / "run.xlsx"
    workbook = Workbook()
    first = workbook.active
    assert first is not None
    for sheet in (first, workbook.create_sheet("Second")):
        sheet.append(["Bill", "Pay"])
        sheet.append(["ML-4471", 44900])
    workbook.save(target)
    workbook.close()
    monkeypatch.setattr(documents, "MAX_WORKBOOK_CELLS", 6)

    result = agent_workspace.read_deliverable(deliverable_root, target)

    assert result.status == "unreadable"


@pytest.mark.parametrize("replacement", ["file", "symlink"])
def test_parser_reads_a_trusted_snapshot_when_the_solver_replaces_its_file(
    deliverable_root: Path, monkeypatch: pytest.MonkeyPatch, replacement: str
):
    target = deliverable_root / "outputs" / "notes.md"
    target.write_text("submitted content", encoding="utf-8")
    secret = deliverable_root.parent / "secret.md"
    secret.write_text("private verifier data", encoding="utf-8")
    parse = agent_workspace.document_text
    observed: list[Path] = []

    def replace_original_before_parsing(snapshot: Path) -> str:
        observed.append(snapshot)
        assert snapshot != target
        assert not snapshot.is_relative_to(deliverable_root)
        assert snapshot.is_file() and not snapshot.is_symlink()
        target.unlink()
        if replacement == "symlink":
            target.symlink_to(secret)
        else:
            target.write_text("replacement content", encoding="utf-8")
        return parse(snapshot)

    monkeypatch.setattr(agent_workspace, "document_text", replace_original_before_parsing)

    result = agent_workspace.read_deliverable(deliverable_root, target)

    assert observed
    assert result.status == "saved"
    assert result.text == "submitted content"
