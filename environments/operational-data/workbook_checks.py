"""Server-side checks of saved workbook values against task-authored reference rows."""

from __future__ import annotations

import math
import re
from collections.abc import Sequence
from datetime import datetime, timedelta
from io import BytesIO
from typing import TYPE_CHECKING, cast

from openpyxl import load_workbook
from pydantic import (
    BaseModel,
    ConfigDict,
    Field,
    StrictFloat,
    StrictInt,
    StrictStr,
    model_validator,
)

if TYPE_CHECKING:
    from openpyxl.workbook.workbook import Workbook

    from agent_workspace import Deliverable

CellValue = StrictStr | StrictInt | StrictFloat | None


class WorkbookTable(BaseModel):
    """An exact keyed table; extra columns and notes outside its key pattern are allowed."""

    model_config = ConfigDict(extra="forbid", allow_inf_nan=False)
    sheet: str
    keys: list[str] = Field(min_length=1)
    rows: list[dict[str, CellValue]] = Field(min_length=1)
    required_columns: list[str] = Field(default_factory=list)
    key_pattern: str | None = None
    utc_columns: list[str] = Field(default_factory=list)
    totals: dict[str, StrictInt | StrictFloat] = Field(default_factory=dict)
    optional_assertions: list[dict[str, CellValue]] = Field(default_factory=list)

    @model_validator(mode="after")
    def validate_rows(self) -> WorkbookTable:
        columns = set(self.rows[0])
        if not set(self.keys) <= columns or any(set(row) != columns for row in self.rows):
            raise ValueError("reference rows must share columns and include every key")
        keys = [tuple(row[key] for key in self.keys) for row in self.rows]
        if len(set(keys)) != len(keys):
            raise ValueError("reference row keys must be unique")
        for assertion in self.optional_assertions:
            if (
                not set(self.keys) <= set(assertion)
                or tuple(assertion[key] for key in self.keys) not in keys
            ):
                raise ValueError("optional assertions must refer to required row keys")
        if not (set(self.utc_columns) | set(self.totals)) <= columns:
            raise ValueError("timestamp and total columns must be present in reference rows")
        if self.key_pattern is not None:
            re.compile(self.key_pattern)
        return self


class WorkbookAlternative(BaseModel):
    """All tables in an alternative must match the same saved workbook."""

    model_config = ConfigDict(extra="forbid")
    tables: list[WorkbookTable] = Field(min_length=1)


class WorkbookCheck(BaseModel):
    model_config = ConfigDict(extra="forbid")
    file: str
    alternatives: list[WorkbookAlternative] = Field(min_length=1)
    tolerance: float = Field(default=0.000001, ge=0, allow_inf_nan=False)


def _equal(actual: object, expected: CellValue, *, tolerance: float, utc: bool) -> bool:
    if isinstance(expected, (int, float)):
        return (
            isinstance(actual, (int, float))
            and not isinstance(actual, bool)
            and math.isfinite(actual)
            and abs(actual - expected) <= tolerance
        )
    if utc and expected is not None:
        if not isinstance(actual, str) or not isinstance(expected, str):
            return False
        try:
            actual_time = datetime.fromisoformat(actual.replace("Z", "+00:00"))
            expected_time = datetime.fromisoformat(expected.replace("Z", "+00:00"))
        except ValueError:
            return False
        return actual_time.utcoffset() == timedelta(0) and actual_time == expected_time
    return actual == expected


def _table_errors(workbook: Workbook, table: WorkbookTable, tolerance: float) -> list[str]:
    if table.sheet not in workbook.sheetnames:
        return [f"Missing worksheet: {table.sheet}"]
    sheet = workbook[table.sheet]
    columns = set(table.rows[0]) | set(table.required_columns)
    headers: list[object] | None = None
    header_row = 0
    for number, row in enumerate(sheet.iter_rows(values_only=True), start=1):
        if columns <= set(row):
            headers, header_row = list(row), number
            break
    if headers is None:
        return [f"{table.sheet}: missing required headers {sorted(columns)}"]
    columns.update(
        column
        for assertion in table.optional_assertions
        for column in assertion
        if column in headers
    )
    if any(headers.count(column) != 1 for column in columns):
        return [f"{table.sheet}: duplicate required headers"]
    positions = {column: headers.index(column) for column in columns}
    expected_rows = {tuple(row[key] for key in table.keys): row for row in table.rows}
    actual_rows: dict[tuple[object, ...], dict[str, object]] = {}
    errors: list[str] = []
    for values in sheet.iter_rows(min_row=header_row + 1, values_only=True):
        row = {
            column: values[index] if index < len(values) else None
            for column, index in positions.items()
        }
        key = tuple(row[column] for column in table.keys)
        if all(value is None for value in key):
            continue
        if table.key_pattern is not None and re.fullmatch(table.key_pattern, str(key[0])) is None:
            continue
        if key in actual_rows:
            errors.append(f"{table.sheet}: duplicate key {key}")
        actual_rows[key] = row
    if set(actual_rows) != set(expected_rows):
        errors.append(f"{table.sheet}: row keys differ from the required set")
    for key, expected in expected_rows.items():
        if key not in actual_rows:
            continue
        for column, value in expected.items():
            actual = actual_rows[key][column]
            if not _equal(actual, value, tolerance=tolerance, utc=column in table.utc_columns):
                errors.append(f"{table.sheet} {key} {column}: expected {value!r}, got {actual!r}")
    for column, expected in table.totals.items():
        numbers = [row[column] for row in actual_rows.values() if row[column] is not None]
        if not all(type(value) in (int, float) for value in numbers) or not _equal(
            sum(cast("list[int | float]", numbers)), expected, tolerance=tolerance, utc=False
        ):
            errors.append(f"{table.sheet}: derived total for {column} differs from {expected}")
    for assertion in table.optional_assertions:
        key = tuple(assertion[column] for column in table.keys)
        actual_row = actual_rows.get(key, {})
        for column, expected in assertion.items():
            actual = actual_row.get(column)
            if isinstance(expected, (int, float)) and isinstance(actual, str):
                text = actual.strip()
                try:
                    actual = float(text.removesuffix("%")) / (100 if text.endswith("%") else 1)
                except ValueError:
                    pass
            if (
                column not in table.keys
                and actual is not None
                and not _equal(
                    actual, expected, tolerance=tolerance, utc=column in table.utc_columns
                )
            ):
                errors.append(
                    f"{table.sheet} {key} {column}: expected {expected!r}, got {actual!r}"
                )
    return errors


def check_workbook(check: WorkbookCheck, deliverables: Sequence[Deliverable]) -> list[str]:
    """Return no errors when any complete reference alternative matches actual bytes."""
    files = [item for item in deliverables if item.path == f"outputs/{check.file}"]
    if len(files) != 1 or files[0].status != "saved" or not files[0].content:
        return [f"Missing or unreadable workbook: {check.file}"]
    # Deliverable parsing already validated these immutable bytes through the same reader.
    workbook = load_workbook(BytesIO(files[0].content), read_only=True, data_only=True)
    try:
        alternatives = [
            [
                error
                for table in alternative.tables
                for error in _table_errors(workbook, table, check.tolerance)
            ]
            for alternative in check.alternatives
        ]
        if any(not errors for errors in alternatives):
            return []
        return min(alternatives, key=len)
    finally:
        workbook.close()
