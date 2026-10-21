"""Bounded document renderings for source evidence and submitted deliverables."""

from __future__ import annotations

import json
import sqlite3
from collections.abc import Callable, Generator, Iterator
from contextlib import ExitStack, closing
from pathlib import Path
from zipfile import ZipFile

from docx import Document
from openpyxl import load_workbook
from pptx import Presentation
from pypdf import PdfReader

MAX_FILE_BYTES = 16 * 1024 * 1024
MAX_EXPANDED_DOCUMENT_BYTES = 64 * 1024 * 1024
MAX_TEXT_BYTES = 256 * 1024
MAX_WORKBOOK_CELLS = 1_000_000


class UnreadableDocument(Exception):
    """A file whose format is unsupported or whose contents could not be parsed."""


def _rows(lines: Iterator[list[str]]) -> Iterator[str]:
    for cells in lines:
        if any(cells):
            yield "\t".join(cells).rstrip("\t")


def _xlsx(path: Path) -> Generator[str, None, None]:
    """JSON rows preserve stored cell types; uncached formulas remain explicit."""
    with ExitStack() as opened:
        formulas = load_workbook(path, read_only=True, data_only=False)
        opened.callback(formulas.close)
        values = load_workbook(path, read_only=True, data_only=True)
        opened.callback(values.close)
        visited_cells = 0
        for name in values.sheetnames:
            sheet = values[name]
            if (
                sheet.max_row is not None
                and sheet.max_column is not None
                and sheet.max_row * sheet.max_column > MAX_WORKBOOK_CELLS
            ):
                raise UnreadableDocument(f"exceeds the {MAX_WORKBOOK_CELLS}-cell workbook limit")
            yield f"## Sheet: {name}"
            rows = zip(
                sheet.iter_rows(values_only=True),
                formulas[name].iter_rows(),
                strict=True,
            )
            for value_row, formula_row in rows:
                visited_cells += max(1, len(value_row))
                if visited_cells > MAX_WORKBOOK_CELLS:
                    raise UnreadableDocument(
                        f"exceeds the {MAX_WORKBOOK_CELLS}-cell workbook limit"
                    )
                cells = [
                    f"[uncached formula: {cell.value}]"
                    if value is None and cell.data_type == "f"
                    else value
                    for value, cell in zip(value_row, formula_row, strict=True)
                ]
                if any(value is not None for value in cells):
                    yield json.dumps(cells, ensure_ascii=False, default=str)


def _docx(path: Path) -> Generator[str, None, None]:
    document = Document(str(path))
    for paragraph in document.paragraphs:
        if paragraph.text.strip():
            yield paragraph.text
    for table in document.tables:
        yield "## Table"
        yield from _rows([cell.text for cell in row.cells] for row in table.rows)


def _pptx(path: Path) -> Generator[str, None, None]:
    for number, slide in enumerate(Presentation(str(path)).slides, start=1):
        yield f"## Slide {number}"
        for shape in slide.shapes:
            if shape.has_text_frame and shape.text_frame.text.strip():
                yield shape.text_frame.text
            if shape.has_table:
                yield from _rows([cell.text for cell in row.cells] for row in shape.table.rows)


def _pdf(path: Path) -> Generator[str, None, None]:
    pages = PdfReader(path).pages
    for number, page in enumerate(pages, 1):
        yield f"## Page {number}\n{page.extract_text() or ''}"


def _sqlite(path: Path) -> Generator[str, None, None]:
    with closing(sqlite3.connect(f"{path.resolve().as_uri()}?mode=ro", uri=True)) as connection:
        tables = connection.execute(
            "SELECT name FROM sqlite_schema "
            "WHERE type = 'table' AND name NOT LIKE 'sqlite_%' ORDER BY name"
        )
        for (name,) in tables:
            identifier = name.replace('"', '""')
            with closing(connection.execute(f'SELECT * FROM "{identifier}"')) as rows:
                yield f"## Table: {name}"
                yield json.dumps([column[0] for column in rows.description], ensure_ascii=False)
                for row in rows:
                    yield json.dumps(row, ensure_ascii=False, default=str)


_READERS: dict[str, Callable[[Path], Generator[str, None, None]]] = {
    ".xlsx": _xlsx,
    ".xlsm": _xlsx,
    ".docx": _docx,
    ".pptx": _pptx,
    ".pdf": _pdf,
    ".sqlite": _sqlite,
    ".sqlite3": _sqlite,
    ".db": _sqlite,
}


def _bounded_text(lines: Iterator[str]) -> str:
    parts: list[str] = []
    size = 0
    for line in lines:
        size += len(line.encode("utf-8")) + bool(parts)
        if size > MAX_TEXT_BYTES:
            raise UnreadableDocument(f"exceeds the {MAX_TEXT_BYTES}-byte text limit")
        parts.append(line)
    return "\n".join(parts)


def document_text(path: Path) -> str:
    """The text of an office document, a PDF or a UTF-8 text file."""
    if path.stat().st_size > MAX_FILE_BYTES:
        raise UnreadableDocument(f"exceeds the {MAX_FILE_BYTES}-byte input limit")
    suffix = path.suffix.lower()
    reader = _READERS.get(suffix)
    if reader is None:
        try:
            return _bounded_text(iter([path.read_text(encoding="utf-8")]))
        except UnicodeDecodeError as error:
            raise UnreadableDocument("not a text file or a supported document format") from error
    try:
        if suffix in {".xlsx", ".xlsm", ".docx", ".pptx"}:
            with ZipFile(path) as archive:
                if sum(item.file_size for item in archive.infolist()) > MAX_EXPANDED_DOCUMENT_BYTES:
                    raise UnreadableDocument(
                        f"exceeds the {MAX_EXPANDED_DOCUMENT_BYTES}-byte expanded document limit"
                    )
        with closing(reader(path)) as lines:
            return _bounded_text(lines)
    except UnreadableDocument:
        raise
    except Exception as error:
        # Agent-written files are untrusted input; any parser failure means unreadable.
        raise UnreadableDocument(f"could not be read as {suffix}: {error}") from error
